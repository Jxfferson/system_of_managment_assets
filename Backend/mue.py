#!/usr/bin/env python3
"""
Generador de Manual de Usuario - Asset Management
Sistema: OTD Support
Autor: Jefferson Stic Correa Lopez
Fecha: Mayo 2026

Instrucciones:
1. Instalar: pip install python-docx
2. Ejecutar: python generate_user_manual.py
3. Archivo generado: Manual_Usuario_AssetManagement_v2.0.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

def add_heading(doc, text, level=1):
    """Agregar encabezado"""
    heading = doc.add_heading(text, level=level)
    heading.runs[0].font.color.rgb = RGBColor(0, 51, 102)
    return heading

def add_table(doc, headers, rows):
    """Agregar tabla con formato"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
    
    for row in rows:
        row_cells = table.add_row().cells
        for i, cell_text in enumerate(row):
            row_cells[i].text = str(cell_text)
    
    return table

def generate_manual():
    """Generar el manual completo"""
    
    doc = Document()
    
    # PORTADA
    title = doc.add_heading('MANUAL DE USUARIO Y ADMINISTRADOR', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Modulo de Asset Management (Gestion de Activos)')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    
    doc.add_paragraph()
    
    # Metadata
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = 'Light Grid'
    meta_data = [
        ('Version', '2.0'),
        ('Fecha', datetime.datetime.now().strftime('%B %Y')),
        ('Sistema', 'OTD Support - Asset Management'),
        ('Autor', 'Jefferson Stic Correa Lopez - OTD TEAM DEVELOPMENT')
    ]
    for i, (key, value) in enumerate(meta_data):
        meta_table.cell(i, 0).text = key
        meta_table.cell(i, 1).text = value
        meta_table.cell(i, 0).paragraphs[0].runs[0].font.bold = True
    
    doc.add_page_break()
    
    # CONTENIDO
    add_heading(doc, 'CONTENIDO', level=1)
    sections = [
        '1. Introduccion',
        '2. Acceso al Sistema',
        '3. Panel Principal y Navegacion',
        '4. Dashboard de Estadisticas',
        '5. Gestion de Activos',
        '6. Funcionalidades Avanzadas',
        '7. Filtrar y Buscar',
        '8. Exportar a Excel',
        '9. Atajos de Teclado',
        'Resumen Rapido'
    ]
    for section in sections:
        doc.add_paragraph(section, style='List Bullet')
    
    doc.add_page_break()
    
    # 1. INTRODUCCION
    add_heading(doc, '1. INTRODUCCION', level=1)
    doc.add_paragraph(
        'Este manual describe el uso del modulo Asset Management, disenado para el control, '
        'seguimiento y analisis de activos tecnologicos (teclados, mouse, cables, monitores, '
        'Ethernet, etc.) asignados a estaciones de trabajo.'
    )
    
    doc.add_paragraph('Funcionalidades principales:', style='Heading 3')
    features = [
        'Registro individual y masivo de activos',
        'Seguimiento por estacion y monitor (Izquierdo/Derecho)',
        'Dashboard de estadisticas y alertas de stock critico',
        'Integracion con escaner de codigos/seriales',
        'Gestion de estaciones y activos asignados',
        'Filtros avanzados y exportacion a Excel'
    ]
    for feature in features:
        doc.add_paragraph(feature, style='List Bullet')
    
    # 2. ACCESO AL SISTEMA
    add_heading(doc, '2. ACCESO AL SISTEMA', level=1)
    add_heading(doc, '2.1 Iniciar Sesion', level=2)
    steps_login = [
        'Haga clic en el boton "Admin" (esquina superior derecha)',
        'Ingrese su contrasena de administrador',
        'Presione "Login"',
        'Accedera al panel principal de Asset Management'
    ]
    for i, step in enumerate(steps_login, 1):
        doc.add_paragraph(f'Paso {i}: {step}')
    
    doc.add_paragraph()
    doc.add_paragraph('Nota: El acceso esta restringido a perfiles con rol de Administrador o Operador de TI.', style='Intense Quote')
    
    doc.add_page_break()
    
    # 3. PANEL PRINCIPAL
    add_heading(doc, '3. PANEL PRINCIPAL Y NAVEGACION', level=1)
    doc.add_paragraph(
        'Al ingresar, encontrara tres pestanas principales en la parte superior izquierda:'
    )
    
    tabs = [
        ('Assets', 'Vista completa de todos los activos registrados'),
        ('Items', 'Catalogo base de tipos de activos disponibles'),
        ('Statistics', 'Dashboard analitico en tiempo real')
    ]
    for tab, desc in tabs:
        doc.add_paragraph(f'{tab}: {desc}', style='List Bullet')
    
    add_heading(doc, '3.1 Tarjetas de Resumen', level=2)
    summary_headers = ['Tarjeta', 'Descripcion']
    summary_rows = [
        ['Total Assets', 'Cantidad total de activos registrados'],
        ['Total Stock', 'Activos disponibles en inventario (verde)'],
        ['Total Exits', 'Activos con fecha de salida (amarillo)'],
        ['Distribution', 'Desglose por tipo: Return, Missing, Damage']
    ]
    add_table(doc, summary_headers, summary_rows)
    
    add_heading(doc, '3.2 Botones de Accion', level=2)
    buttons_headers = ['Boton', 'Funcion']
    buttons_rows = [
        ['+ Add Asset', 'Registrar activo individual'],
        ['+ Add Lot', 'Registrar lote masivo (5-999 activos)'],
        ['Filter', 'Abrir panel de filtros avanzados'],
        ['Export', 'Descargar tabla actual a Excel'],
        ['Admin', 'Menu: Change Password, Logout'],
        ['Scanner', 'Configurar y activar escaner']
    ]
    add_table(doc, buttons_headers, buttons_rows)
    
    add_heading(doc, '3.3 Tabla de Activos - Columnas', level=2)
    doc.add_paragraph('ITEM | SERIAL | ENTRY DATE | EXIT DATE | RETURN TYPE | OBSERVATIONS | DESTINATION | MONITOR | CURRENT HEADQUARTERS')
    
    doc.add_page_break()
    
    # 4. DASHBOARD ESTADISTICAS
    add_heading(doc, '4. DASHBOARD DE ESTADISTICAS', level=1)
    doc.add_paragraph('Acceda haciendo clic en la pestana "Statistics".')
    
    add_heading(doc, '4.1 Active Alerts', level=2)
    doc.add_paragraph(
        'Muestra items con stock critico que requieren atencion inmediata. '
        'Indicador "Critical" cuando stock disponible = 0 o muy bajo. '
        'Boton "Create Order" para solicitar reposicion.'
    )
    
    add_heading(doc, '4.2 Item Availability', level=2)
    doc.add_paragraph('Tabla detallada por tipo de activo:')
    availability_headers = ['Campo', 'Descripcion']
    availability_rows = [
        ['Available', 'Unidades libres en inventario'],
        ['Total', 'Unidades registradas en el sistema'],
        ['% Status', 'Barra visual (rojo <30%, verde >70%)']
    ]
    add_table(doc, availability_headers, availability_rows)
    
    add_heading(doc, '4.3 Resumen Inferior', level=2)
    doc.add_paragraph('TOTAL UNITS: Total general de activos')
    doc.add_paragraph('AVAILABLE: Disponibles para asignacion')
    doc.add_paragraph('ASSIGNED: Actualmente en estaciones o en uso')
    
    doc.add_page_break()
    
    # 5. GESTION DE ACTIVOS
    add_heading(doc, '5. GESTION DE ACTIVOS', level=1)
    
    add_heading(doc, '5.1 Agregar Activo Individual', level=2)
    doc.add_paragraph('Pasos:')
    steps_add = [
        'Clic en "+ Add Asset"',
        'Complete el formulario (ver tabla de campos)',
        'Clic en "Save"'
    ]
    for i, step in enumerate(steps_add, 1):
        doc.add_paragraph(f'Paso {i}: {step}', style='List Number')
    
    add_heading(doc, 'Campos del Formulario', level=3)
    form_headers = ['Campo', 'Descripcion', 'Obligatorio']
    form_rows = [
        ['Select or type...', 'Seleccione del catalogo o escriba nombre nuevo', 'Si'],
        ['Serial Number', 'Autogenerado o ingresado manualmente', 'Si'],
        ['Entry Date', 'Fecha de ingreso (mm/dd/yyyy)', 'Si'],
        ['Exit Date', 'Fecha de salida (si aplica)', 'No'],
        ['Return Type', 'Damage / Missing / Return', 'No'],
        ['Observations', 'Notas sobre el estado o motivo', 'No'],
        ['Destination', 'Estacion o ubicacion final', 'Si (si se asigna)'],
        ['Monitor', 'Left / Right / N/A', 'No']
    ]
    add_table(doc, form_headers, form_rows)
    
    add_heading(doc, '5.2 Editar Activo', level=2)
    edit_steps = [
        'Localice el activo en la tabla',
        'Clic en el icono de lapiz (editar) en la columna Actions',
        'Modifique los campos permitidos',
        'Clic en "Save"'
    ]
    for i, step in enumerate(edit_steps, 1):
        doc.add_paragraph(f'Paso {i}: {step}', style='List Number')
    
    doc.add_paragraph('Campos editables: Item, Entry/Exit Date, Destination, Return Type, Observations, Monitor. El Serial no es editable.')
    
    add_heading(doc, '5.3 Eliminar Activo', level=2)
    delete_steps = [
        'Clic en el icono de basura',
        'Confirme la eliminacion en el dialogo emergente'
    ]
    for i, step in enumerate(delete_steps, 1):
        doc.add_paragraph(f'Paso {i}: {step}', style='List Number')
    
    doc.add_paragraph('Importante: La accion es permanente. Se recomienda exportar respaldo antes.', style='Intense Quote')
    
    doc.add_page_break()
    
    # 6. FUNCIONALIDADES AVANZADAS
    add_heading(doc, '6. FUNCIONALIDADES AVANZADAS', level=1)
    
    add_heading(doc, '6.1 Integracion con Escaner', level=2)
    doc.add_paragraph('Pasos:')
    scanner_steps = [
        'Haga clic en el boton "Scanner" (junto a + Add Asset)',
        'Seleccione una opcion:',
        '  - Configure Scanner: Vincular dispositivo USB/Bluetooth',
        '  - Activate Scanner: Encender lectura en tiempo real',
        '  - Scanning Mode: Cambiar entre Register, Check Out, Return',
        '  - Scanner Help: Guia rapida de uso'
    ]
    for step in scanner_steps:
        doc.add_paragraph(step, style='List Bullet')
    
    add_heading(doc, '6.2 Modal: Station Assets', level=2)
    doc.add_paragraph(
        'Al hacer clic en una estacion o ubicacion, se abre un panel flotante que muestra '
        'los activos asignados a esa estacion. Botones por item: Move (trasladar) y '
        'Unassign (devolver a inventario).'
    )
    
    add_heading(doc, '6.3 Agregar Lote Masivo', level=2)
    lot_steps = [
        'Clic en "+ Add Lot"',
        'Seleccione el tipo de activo base',
        'Ingrese la cantidad (5-999)',
        'El sistema generara seriales consecutivos automaticamente',
        'Revise y confirme con "Save Batch"'
    ]
    for i, step in enumerate(lot_steps, 1):
        doc.add_paragraph(f'Paso {i}: {step}', style='List Number')
    
    doc.add_page_break()
    
    # 7. FILTRAR Y BUSCAR
    add_heading(doc, '7. FILTRAR Y BUSCAR', level=1)
    
    add_heading(doc, '7.1 Abrir Filtros', level=2)
    doc.add_paragraph('Clic en "Filter" (esquina superior derecha).')
    
    add_heading(doc, '7.2 Opciones Disponibles', level=2)
    filter_headers = ['Filtro', 'Uso']
    filter_rows = [
        ['Filter by Name', 'Busqueda parcial por nombre del activo'],
        ['Filter by Serial', 'Busqueda exacta o parcial por serial'],
        ['Destination', 'Filtrar por estacion/ubicacion asignada'],
        ['Entry/Exit Date', 'Rango de fechas de ingreso o salida'],
        ['Return Type', 'Damage / Missing / Return / All'],
        ['Observations', 'Busqueda por palabras clave en notas'],
        ['Monitor', 'Left / Right / N/A']
    ]
    add_table(doc, filter_headers, filter_rows)
    
    add_heading(doc, '7.3 Consejos', level=2)
    tips = [
        'Puede combinar multiples filtros',
        'Las busquedas son parciales (no requiere coincidencia exacta)',
        'Para limpiar: borre los campos o recargue la vista'
    ]
    for tip in tips:
        doc.add_paragraph(tip, style='List Bullet')
    
    doc.add_page_break()
    
    # 8. EXPORTAR A EXCEL
    add_heading(doc, '8. EXPORTAR A EXCEL', level=1)
    
    add_heading(doc, '8.1 Exportar Todo o Filtrado', level=2)
    export_steps = [
        'Aplique filtros si desea un subconjunto (opcional)',
        'Clic en "Export"',
        'Seleccione "Export as Excel"',
        'El archivo .xlsx se descargara automaticamente'
    ]
    for i, step in enumerate(export_steps, 1):
        doc.add_paragraph(f'Paso {i}: {step}', style='List Number')
    
    add_heading(doc, '8.2 Columnas Exportadas', level=2)
    doc.add_paragraph('ITEM | SERIAL | ENTRY DATE | EXIT DATE | RETURN TYPE | OBSERVATIONS | DESTINATION | MONITOR | CURRENT HEADQUARTERS')
    
    add_heading(doc, '8.3 Usos Recomendados', level=2)
    uses = [
        'Respaldos periodicos del inventario',
        'Reportes gerenciales o auditorias',
        'Analisis avanzado en Excel (tablas dinamicas, graficos)'
    ]
    for use in uses:
        doc.add_paragraph(use, style='List Bullet')
    
    doc.add_page_break()
    
    # 9. ATAJOS DE TECLADO
    add_heading(doc, '9. ATAJOS DE TECLADO', level=1)
    
    shortcut_headers = ['Atajo', 'Accion', 'Descripcion']
    shortcut_rows = [
        ['Ctrl + Alt + N', 'Nuevo Activo', 'Abre formulario de registro individual'],
        ['Ctrl + Alt + L', 'Nuevo Lote', 'Abre formulario de carga masiva'],
        ['Ctrl + Alt + E', 'Exportar Excel', 'Descarga tabla actual (con/sin filtros)'],
        ['Ctrl + Alt + C', 'Limpiar Filtros', 'Restablece la vista completa'],
        ['Ctrl + Alt + S', 'Guardar Cambios', 'Confirma creacion o edicion']
    ]
    add_table(doc, shortcut_headers, shortcut_rows)
    
    doc.add_paragraph()
    doc.add_paragraph('Nota para Mac: Reemplace Ctrl por Cmd (⌘)', style='Intense Quote')
    doc.add_paragraph('Los atajos no interfieren con la escritura en campos de texto.')
    doc.add_paragraph('Al pasar el cursor sobre botones principales, se muestra su atajo correspondiente.')
    
    doc.add_page_break()
    
    # RESUMEN RAPIDO
    add_heading(doc, 'RESUMEN RAPIDO', level=1)
    
    summary_headers = ['Accion', 'Pasos']
    summary_rows = [
        ['Agregar activo', '+ Add Asset -> Formulario -> Save'],
        ['Agregar lote', '+ Add Lot -> Cantidad -> Save Batch'],
        ['Editar activo', 'Icono editar -> Modificar -> Save'],
        ['Eliminar activo', 'Icono basura -> Confirmar'],
        ['Ver activos por estacion', 'Clic en estacion -> Modal Station Assets'],
        ['Usar escaner', 'Boton Scanner -> Activar -> Seleccionar modo'],
        ['Filtrar', 'Filter -> Seleccionar criterios -> Aplicar'],
        ['Exportar', 'Export -> Export as Excel'],
        ['Ver estadisticas', 'Pestana Statistics -> Alertas y disponibilidad']
    ]
    add_table(doc, summary_headers, summary_rows)
    
    # FOOTER
    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(
        f'\nVersion: 2.0\n'
        f'Fecha: {datetime.datetime.now().strftime("%B %Y")}\n'
        f'Sistema: OTD Support - Asset Management\n'
        f'Elaborado por: Jefferson Stic Correa Lopez - OTD TEAM DEVELOPMENT'
    )
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(102, 102, 102)
    
    # Guardar
    output_filename = 'Manual_Usuario_AssetManagement_v2.0.docx'
    doc.save(output_filename)
    print(f'Documento generado: {output_filename}')
    return output_filename

if __name__ == '__main__':
    try:
        generate_manual()
        print('\nListo! Abra el archivo .docx con Microsoft Word o Google Docs')
    except ImportError:
        print('Error: Instale python-docx con: pip install python-docx')
    except Exception as e:
        print(f'Error: {e}')
        raise